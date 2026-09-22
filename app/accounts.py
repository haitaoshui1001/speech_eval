"""账号录入：管理员建号与 Excel 批量导入的模板生成、解析与校验。

站点不开放自助注册，账号只能由管理员在本系统里开通，途径有二：
  - 后台单个填写（main.py 的 /admin/users/create）
  - 下载模板 → Excel 填好 → 上传导入（/admin/users/template、/admin/users/import）

Excel 只用 openpyxl 读写 .xlsx（本环境 pandas 与 NumPy 2 存在 ABI 冲突，不可依赖）。
`plan_users()` 是纯函数：输入二维表格数据 + 已存在账号名，输出逐行判定结果，
不碰数据库也不碰文件系统，因此可被 tests/selfcheck.py 直接断言。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from typing import Any, Iterable, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .security import password_ok, valid_username

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEMPLATE_NAME = "账号导入模板.xlsx"
SHEET = "账号"
SHEET_HELP = "填写说明"
HEADERS: tuple[str, ...] = ("用户名", "登录口令", "显示名")

MAX_BYTES = 4 * 1024 * 1024
MAX_ROWS = 5000
COMMENT_PREFIX = "#"

EXAMPLES: tuple[tuple[str, str, str], ...] = (
    ("#zhang01", "k9tq2m", "张强 · 示例行，导入时自动跳过"),
    ("#li02", "p3w7xr", "李娜 · 示例行，导入时自动跳过"),
)

NOTES: tuple[str, ...] = (
    "第 1 行是表头，不要修改、删除或调换顺序；多出来的列会被忽略。",
    "用户名：3–32 位，只能用字母、数字和 . _ -，全站唯一，用户凭它登录。",
    "登录口令：至少 6 位。导入后系统只存加盐哈希，任何页面都不会再显示原文。",
    "显示名：选填，留空时与用户名相同，只作为页面和报告里的称呼。",
    f"用户名以 {COMMENT_PREFIX} 开头的行按示例/注释处理，导入时跳过。",
    "导入只新增账号，不会删除或改动已有账号；重名、缺项的行会带行号退回。",
    f"一次导入建议不超过 {MAX_ROWS} 行，文件不超过 {MAX_BYTES // 1024 // 1024} MB。",
    "口令属于敏感信息：Excel 用完请删除，不要转发明文口令。",
)


class AccountError(Exception):
    """模板/文件层面的错误：整份表格无法处理时抛出。"""


@dataclass
class Row:
    """一行账号数据及其判定结果。status: ready / error / skip / created。"""

    line: int
    username: str = ""
    password: str = ""
    display_name: str = ""
    status: str = "error"
    errors: list[str] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return self.status == "ready"

    @property
    def reason(self) -> str:
        return "；".join(self.errors)

    def reject(self, why: str) -> None:
        self.status = "error"
        self.errors.append(why)

    def to_dict(self) -> dict[str, Any]:
        return {"line": self.line, "username": self.username, "display_name": self.display_name,
                "ok": self.status != "error", "created": self.status == "created",
                "reason": self.reason or ("已跳过" if self.status == "skip" else "")}


def cell_text(value: Any) -> str:
    """Excel 单元格 → 干净的字符串。

    数字型单元格（学号常被存成数字）去掉 `.0` 尾巴，避免 20250101 变成 20250101.0。
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def is_header(cells: Sequence[str]) -> bool:
    return bool(cells) and cells[0] in {HEADERS[0], HEADERS[0].lower(), "username", "用户名".lower()}


def plan_users(table: Iterable[Sequence[Any]], existing: Iterable[str] = ()) -> list[Row]:
    """把表格数据逐行判定成可导入 / 需退回的账号行。

    - 第一条非空行若长得像表头，当作表头丢弃（模板下载后原样上传也能用）。
    - 用户名以 # 开头 → skip（示例与注释）。
    - 三列全空 → 直接忽略，不计入任何统计。
    - 重名分两种：库里已有、本文件内更早的行已有，都退回并写明冲突行号。
    """
    known = {str(u).strip() for u in existing}
    rows: list[Row] = []
    seen: dict[str, int] = {}
    at_top = True
    for line, raw in enumerate(table, start=1):
        cells = [cell_text(v) for v in (raw or ())][:len(HEADERS)]
        cells += [""] * (len(HEADERS) - len(cells))
        username, password, display_name = cells
        if not any(cells):
            continue
        row = Row(line=line, username=username, password=password, display_name=display_name)
        top, at_top = at_top, False
        if top and is_header(cells):
            continue
        if username.startswith(COMMENT_PREFIX):
            row.status = "skip"
            rows.append(row)
            continue
        if is_header(cells):
            row.status = "skip"
            rows.append(row)
            continue
        if not username:
            row.reject("缺用户名")
        elif not valid_username(username):
            row.reject("用户名需 3–32 位，仅限字母、数字与 . _ -")
        if not password:
            row.reject("缺登录口令")
        elif not password_ok(password):
            row.reject("登录口令至少 6 位")
        if username and not row.errors:
            if username in known:
                row.reject("该账号已存在，未做任何改动")
            elif username in seen:
                row.reject(f"与模板第 {seen[username]} 行重名")
            else:
                seen[username] = line
                row.status = "ready"
        rows.append(row)
    return rows


def summarize(rows: Sequence[Row], filename: str = "") -> dict[str, Any]:
    """把逐行结果整理成页面回执：成功清单、退回清单、跳过数。"""
    created = [r for r in rows if r.status == "created"]
    failed = [r for r in rows if r.status == "error"]
    skipped = [r for r in rows if r.status == "skip"]
    return {
        "filename": filename,
        "created": [r.username for r in created],
        "failed": [r.to_dict() for r in failed],
        "rows": [r.to_dict() for r in rows if r.status != "skip"],
        "skipped": len(skipped),
        "ok": not failed,
    }


def build_template() -> bytes:
    """生成 .xlsx 模板：一个填写用的「账号」表 + 一个「填写说明」表。"""
    head_font = Font(bold=True, color="FFFFFF", size=11)
    head_fill = PatternFill("solid", fgColor="2F4858")
    demo_font = Font(italic=True, color="8A8A8A", size=11)

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(list(HEADERS))
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = head_font
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for demo in EXAMPLES:
        ws.append(list(demo))
        for col in range(1, len(HEADERS) + 1):
            ws.cell(row=ws.max_row, column=col).font = demo_font
    for col, width in enumerate((20, 18, 34), start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    help_sheet = wb.create_sheet(SHEET_HELP)
    help_sheet.append([f"{TEMPLATE_NAME} 填写说明"])
    help_sheet.cell(row=1, column=1).font = Font(bold=True, size=12)
    for note in NOTES:
        help_sheet.append([note])
    help_sheet.cell(row=help_sheet.max_row + 2, column=1).value = "列对照：" + " ｜ ".join(HEADERS)
    help_sheet.column_dimensions["A"].width = 92
    for r in range(2, help_sheet.max_row + 1):
        help_sheet.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def read_table(data: bytes, filename: str = "") -> list[Sequence[Any]]:
    """xlsx 字节 → 二维表格数据；文件层面不合法就抛 AccountError。"""
    suffix = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in (filename or "") else ""
    if suffix and suffix != ".xlsx":
        raise AccountError(f"只支持 .xlsx（当前是 {suffix}）。请在 Excel 里「另存为 → xlsx」后重传")
    if not data:
        raise AccountError("上传的文件是空的")
    if len(data) > MAX_BYTES:
        raise AccountError(f"文件超过 {MAX_BYTES // 1024 // 1024} MB 上限")
    try:
        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise AccountError("读不了这个文件：请用模板页「下载 Excel 模板」另存后填写再传") from None
    try:
        sheet = wb[SHEET] if SHEET in wb.sheetnames else wb[wb.sheetnames[0]]
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    finally:
        wb.close()
    if len(rows) > MAX_ROWS:
        raise AccountError(f"表格 {len(rows)} 行，超过 {MAX_ROWS} 行上限，请分次导入")
    if not any(any(cell_text(v) for v in (r or ())) for r in rows):
        raise AccountError("表格里没有内容")
    return rows
